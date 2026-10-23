from datetime import datetime, timedelta, timezone
from html.parser import HTMLParser
from io import BytesIO
from math import isfinite

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models.database import listar_ordens_trabalho, obter_sincronizacao
from services.pcm_planning import ConflictError, text


NF_SECTOR = "RES-TC-NF"
# Physical column positions and merges from the supplied Report worksheet.
HEADERS = {
    "A": "Ordem de Trabalho", "B": "Descrição", "D": "Local", "E": "Status da OT",
    "F": "Tipo de trabalho", "G": "Refere-se a", "I": "Setor de manutenção",
    "J": "Grupo de Preparadores", "K": "Data de notificação", "M": "Não iniciar antes de",
    "R": "Terminar antes de", "T": "Data de vencimento do PM",
    "U": "Horas de mão de obra estimadas", "V": "Duração", "W": "Início planejado",
    "Y": "Final planejado", "AA": "PM", "AB": "Plano de Trabalho", "AC": "Ldtext",
}
WIDTHS = [2340, 3876, 1389, 3510, 3510, 2450, 365, 2121, 3510, 3510,
          2450, 2450, 1755, 365, 365, 1060, 1389, 2450, 2450, 2450,
          2450, 2450, 2450, 2450, 2450, 2450, 5266, 3510, 5266]


class PlainText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)

    def handle_starttag(self, tag, attrs):
        if tag in {"br", "p", "div", "li"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"p", "div", "li"}:
            self.parts.append("\n")


def plain_text(value: str) -> str:
    parser = PlainText()
    parser.feed(value)
    return ILLEGAL_CHARACTERS_RE.sub("", "".join(parser.parts).strip())


def write_text(sheet, address, value):
    cell = sheet[address]
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell.data_type = "s"


def write_datetime(sheet, row, date_col, time_col, value):
    if not value:
        return
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"Data BMA inválida na linha {row}, coluna {date_col}.") from exc
    if parsed.tzinfo:
        parsed = parsed.astimezone(timezone(timedelta(hours=-3)))
    sheet[f"{date_col}{row}"] = parsed.date()
    sheet[f"{date_col}{row}"].number_format = "dd/mm/yyyy"
    if time_col and len(value) > 10:
        sheet[f"{time_col}{row}"] = parsed.time().replace(tzinfo=None)
        sheet[f"{time_col}{row}"].number_format = "hh:mm:ss"


def render_report(orders: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.merge_cells("A1:G1")
    sheet.merge_cells("H1:P1")
    sheet.merge_cells("A2:AC2")
    sheet["A2"] = "Planejamento NFE"
    sheet["A2"].font = Font(name="serif", size=20, bold=True, color="333399")
    sheet.row_dimensions[1].height = 15
    sheet.row_dimensions[2].height = 28.5
    sheet.row_dimensions[3].height = 64.5
    for index, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width / 256
    for span in ("B3:C3", "G3:H3", "K3:L3", "M3:Q3", "R3:S3", "W3:X3", "Y3:Z3"):
        sheet.merge_cells(span)
    for column, label in HEADERS.items():
        sheet[f"{column}3"] = label
    for cell in sheet[3]:
        cell.font = Font(name="serif", size=10, bold=True, color="333399")
        cell.fill = PatternFill("solid", fgColor="C0C0C0")
        cell.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
    border = Border(bottom=Side(style="dotted", color="808080"),
                    right=Side(style="hair", color="D9D9D9"))
    for row, order in enumerate(orders, 4):
        for column in range(1, 30):
            cell = sheet.cell(row, column)
            cell.font = Font(name="serif", size=10)
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            cell.border = border
        for left, right in (("B", "C"), ("G", "H"), ("M", "O"), ("P", "Q")):
            sheet.merge_cells(f"{left}{row}:{right}{row}")
        for column, field in {
            "A": "wonum", "B": "description", "D": "location", "E": "status",
            "F": "worktype", "G": "origin", "I": "maintenanceshop", "J": "planner",
            "AA": "pmnum", "AB": "jpnum",
        }.items():
            write_text(sheet, f"{column}{row}", text(order, field))
        write_text(sheet, f"AC{row}", plain_text(text(order, "description_longdescription")))
        for date_col, time_col, field in [
            ("K", "L", "reportdate"), ("M", "P", "sneconstraint"),
            ("R", "S", "fnlconstraint"), ("T", None, "pmduedate"),
            ("W", "X", "schedstart"), ("Y", "Z", "schedfinish"),
        ]:
            write_datetime(sheet, row, date_col, time_col, text(order, field))
        for column, field in (("U", "estlabhrs"), ("V", "estdur")):
            value = text(order, field)
            if value:
                try:
                    hours = float(value)
                except ValueError as exc:
                    raise ValueError(f"Horas BMA inválidas na linha {row}, campo {field}.") from exc
                if not isfinite(hours) or hours < 0:
                    raise ValueError(f"Horas BMA inválidas na linha {row}, campo {field}.")
                sheet[f"{column}{row}"] = hours / 24
                sheet[f"{column}{row}"].number_format = "[h]:mm"
        description_lines = max(1, len(text(order, "description")) // 22 + 1)
        long_lines = max(1, len(sheet[f"AC{row}"].value) // 28 + 1)
        sheet.row_dimensions[row].height = min(409, max(30, 12 * max(description_lines, long_lines)))
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:AC{max(3, sheet.max_row)}"
    sheet.print_title_rows = "1:3"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_synced_report(job_id: str) -> bytes:
    run = obter_sincronizacao(job_id)
    if run is None:
        raise LookupError("Atualização BMA não encontrada.")
    if run["tipo"] != "ots" or run["setores"] != [NF_SECTOR]:
        raise ValueError("O relatório exige uma atualização exclusiva de OTs do setor RES-TC-NF.")
    if run["status"] != "completed":
        raise ConflictError("A atualização do BMA precisa concluir com sucesso antes de gerar o Excel.")
    orders = [order for order in listar_ordens_trabalho([NF_SECTOR])
              if run["iniciada_em"] <= order["sincronizado_em"] <= run["concluida_em"]]
    if len(orders) != run["ots_salvas"]:
        raise ConflictError("A carteira mudou após esta coleta. Atualize o BMA novamente para exportar.")
    return render_report(orders)
