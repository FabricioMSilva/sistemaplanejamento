import copy
import json
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from typing import Literal
from uuid import uuid4

from pydantic import BaseModel, ConfigDict, Field

from models.database import get_connection, listar_mao_de_obra, listar_ordens_trabalho, obter_equipe
from services.labor_planning import PlanejamentoEquipe, calcular_equipe
from services.pcm_capacity import week_start, weekly_totals

TERMINAL = {"CAN", "CLOSE", "COMP", "HISTEDIT"}


class Input(BaseModel):
    model_config = ConfigDict(extra="forbid")


class Requirement(Input):
    duracao: float | None = Field(default=None, gt=0, le=1000, allow_inf_nan=False)
    pessoas: int | None = Field(default=None, ge=1, le=100)
    orgid: str = Field(default="", max_length=100)
    funcao: str = Field(default="", max_length=100)
    todas_especialidades: bool = False
    selecionada: bool = False
    prioridade: int = Field(default=3, ge=1, le=5)
    liberada: bool = False
    observacao: str = Field(default="", max_length=1000)


class Block(Input):
    id: str = Field(min_length=1, max_length=100)
    ot: str = Field(min_length=1, max_length=100)
    dia: date
    inicio: float = Field(default=0, ge=0, le=1000, allow_inf_nan=False)
    horas: float = Field(gt=0, le=1000, allow_inf_nan=False)
    executantes: list[str] = Field(min_length=1, max_length=100)


class Unavailable(Input):
    executante: str = Field(min_length=1, max_length=100)
    dia: date
    motivo: str = Field(min_length=1, max_length=500)


class Placement(Input):
    ot: str = Field(min_length=1, max_length=100)
    executante: str = Field(min_length=1, max_length=100)
    dia: date


class PlanChange(Input):
    revisao: int = Field(ge=1)
    requisitos: dict[str, Requirement]
    blocos: list[Block] = Field(max_length=10000)
    indisponibilidades: list[Unavailable] = Field(default_factory=list, max_length=10000)
    motivo: str = Field(min_length=1, max_length=1000)
    confirmar: bool = False
    atualizar_especialidades: bool = False
    posicionamentos: list[Placement] = Field(default_factory=list, max_length=20000)


class PlanTeamChange(Input):
    revisao: int = Field(ge=1)
    equipe: PlanejamentoEquipe


class NewPlan(Input):
    setores: list[str] = Field(min_length=1, max_length=28)


class Suggestion(Input):
    modo: Literal["pm", "corretiva"]
    alteracoes: PlanChange
    ordens: list[str] | None = Field(default=None, max_length=20000)


class ExportOrders(Input):
    setores: list[str] = Field(min_length=1, max_length=28)
    inicio: date | None = None
    fim: date | None = None
    modo: Literal["pm", "corretiva"]
    ordens: list[str] = Field(max_length=20000)


class ConflictError(ValueError):
    pass


def specialty_matches(person: dict, requirement: dict) -> bool:
    return requirement.get("todas_especialidades", False) or (
        person["orgid"] == requirement["orgid"] and person["funcao"] == requirement["funcao"]
    )


def specialty_selected(requirement: dict) -> bool:
    return requirement.get("todas_especialidades", False) or bool(requirement["orgid"] and requirement["funcao"])


def text(order: dict, key: str) -> str:
    value = order.get("dados", {}).get(key, order.get(key))
    if isinstance(value, dict):
        value = value.get("value", value.get("_value", value.get("description")))
    return str(value) if value is not None else ""


def order_key(order: dict) -> str:
    return f"{order['siteid']}:{order['wonum']}"


def day_value(value: str) -> str | None:
    if not value:
        return None
    try:
        if len(value) == 10:
            return date.fromisoformat(value).isoformat()
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            return parsed.date().isoformat()
        return parsed.astimezone(timezone(timedelta(hours=-3))).date().isoformat()
    except ValueError:
        return None


def in_period(order: dict, start: str, end: str) -> bool:
    values = [day_value(text(order, key)) for key in
              ("targstartdate", "targcompdate", "schedstart", "schedfinish")]
    return any(value and start <= value <= end for value in values) or any(
        a and b and a <= b and a <= end and b >= start
        for a, b in ((values[0], values[1]), (values[2], values[3]))
    )


def mode_for(order: dict) -> str | None:
    kind = text(order, "worktype").strip().upper()
    if kind == "E":
        return "corretiva"
    if kind in ("I", "S") and text(order, "description").strip().upper().startswith("PM"):
        return "pm"
    return None


def eligible(sectors: list[str], start: str, end: str) -> list[dict]:
    return [order for order in listar_ordens_trabalho(sectors)
            if order["status"].upper() not in TERMINAL and
            (mode_for(order) == "corretiva" or (mode_for(order) == "pm" and in_period(order, start, end)))]


def get_plan(plan_id: str, revision: int | None = None) -> dict:
    with get_connection() as conn:
        row = conn.execute("SELECT dados_json FROM pcm_planos WHERE id=?", (plan_id,)).fetchone() if revision is None else conn.execute(
            "SELECT dados_json FROM pcm_revisoes WHERE plano_id=? AND revisao=?", (plan_id, revision)).fetchone()
    if row is None:
        raise LookupError("Planejamento ou revisao nao encontrado.")
    return json.loads(row["dados_json"])


def list_plans() -> list[dict]:
    with get_connection() as conn:
        rows = conn.execute("SELECT dados_json FROM pcm_planos ORDER BY atualizado_em DESC").fetchall()
    return [{key: doc[key] for key in
             ("id", "revisao", "status", "inicio", "fim", "setores", "planejador", "atualizado_em", "motivo")}
            for doc in (json.loads(row["dados_json"]) for row in rows)]


def delete_plan(plan_id: str) -> dict:
    with get_connection() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT id FROM pcm_planos WHERE id=?", (plan_id,)).fetchone()
        if row is None:
            raise LookupError("Planejamento nao encontrado.")
        conn.execute("DELETE FROM pcm_revisoes WHERE plano_id=?", (plan_id,))
        conn.execute("DELETE FROM pcm_planos WHERE id=?", (plan_id,))
    return {"id": plan_id, "deleted": True}


def revisions(plan_id: str) -> list[dict]:
    get_plan(plan_id)
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT revisao,salvo_em,motivo FROM pcm_revisoes WHERE plano_id=? ORDER BY revisao DESC",
            (plan_id,)).fetchall()
    return [dict(row) for row in rows]


def update_plan_team(plan_id: str, change: PlanTeamChange) -> dict:
    current = get_plan(plan_id)
    if current["revisao"] != change.revisao:
        raise ConflictError("O plano foi alterado. Recarregue antes de salvar a equipe.")
    team = change.equipe
    if (sorted(set(sector.upper() for sector in team.setores)) != current["setores"]
            or str(team.inicio) != current["inicio"] or str(team.fim) != current["fim"]):
        raise ValueError("A equipe deve manter os setores e o período deste planejamento.")
    if not team.executantes or not team.planejador.strip():
        raise ValueError("Selecione ao menos um executante e informe o planejador.")
    if any(member.capacidade != "dias_8h" for member in team.executantes):
        raise ValueError("Revise os dias disponíveis de 8 h dos executantes antes de salvar.")

    labor = listar_mao_de_obra(current["setores"])
    calcular_equipe(team, labor)
    labor_people = {(person["orgid"], person["matricula"]): person for person in labor["profissionais"]}
    people = []
    for member in team.executantes:
        person = labor_people[(member.orgid, member.matricula)]
        craft = next(craft for craft in person["especialidades"] if craft["codigo"] == member.funcao)
        people.append({"id": f"{member.orgid}:{member.matricula}", "nome": person["nome"],
                       "orgid": member.orgid, "matricula": member.matricula, "funcao": member.funcao,
                       "funcao_descricao": craft["descricao"],
                       "dias": sorted(str(day) for day in member.dias_disponiveis)})

    doc = copy.deepcopy(current)
    doc["equipe"] = team.model_dump(mode="json")
    doc["pessoas"] = people
    person_ids = {person["id"] for person in people}
    if (any(member not in person_ids for block in doc["blocos"] for member in block["executantes"])
            or any(item["executante"] not in person_ids for item in doc["indisponibilidades"])
            or any(item["executante"] not in person_ids for item in doc.get("posicionamentos", []))):
        raise ValueError("A equipe não pode remover executantes usados em blocos, indisponibilidades ou posicionamentos.")
    validate_plan(doc, current["status"] == "confirmado")
    doc.update(revisao=current["revisao"] + 1, atualizado_em=datetime.now(timezone.utc).isoformat(),
               motivo="Equipe do planejamento atualizada")
    encoded = json.dumps(doc, ensure_ascii=False)
    pdf = schedule_pdf(doc)
    with get_connection() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT revisao,status FROM pcm_planos WHERE id=?", (plan_id,)).fetchone()
        if row is None:
            raise LookupError("Planejamento nao encontrado.")
        if row["revisao"] != change.revisao:
            raise ConflictError("O plano foi alterado. Recarregue antes de salvar a equipe.")
        if doc["status"] == "confirmado":
            validate_confirmed_capacity(conn, plan_id, doc)
        updated = conn.execute(
            "UPDATE pcm_planos SET revisao=?,atualizado_em=?,dados_json=? WHERE id=? AND revisao=?",
            (doc["revisao"], doc["atualizado_em"], encoded, plan_id, change.revisao))
        if updated.rowcount != 1:
            raise ConflictError("O plano foi alterado. Recarregue antes de salvar a equipe.")
        conn.execute("INSERT INTO pcm_revisoes VALUES (?,?,?,?,?,?)",
                     (plan_id, doc["revisao"], doc["atualizado_em"], doc["motivo"], encoded, pdf))
    return doc


def create_plan(sectors: list[str]) -> dict:
    saved = obter_equipe(sectors)
    if saved is None:
        raise ValueError("Defina e salve a mao de obra primeiro.")
    team = PlanejamentoEquipe.model_validate(saved)
    labor = listar_mao_de_obra(sectors)
    calcular_equipe(team, labor)
    if not team.executantes or not team.planejador.strip():
        raise ValueError("Salve uma equipe e identifique o planejador.")
    if any(member.capacidade != "dias_8h" for member in team.executantes):
        raise ValueError("Revise a equipe no modal e salve os dias disponiveis de 8 h antes de alocar.")
    scope = json.dumps([sorted(set(s.upper() for s in sectors)), str(team.inicio), str(team.fim)])
    with get_connection() as conn:
        existing = conn.execute("SELECT id FROM pcm_planos WHERE escopo=?", (scope,)).fetchone()
    if existing:
        return get_plan(existing["id"])
    people = {(p["orgid"], p["matricula"]): p for p in labor["profissionais"]}
    orders = eligible(sectors, str(team.inicio), str(team.fim))
    requirements = {}
    for order in orders:
        raw = text(order, "estdur")
        try:
            duration = float(raw)
        except ValueError:
            duration = 0
        requirements[order_key(order)] = Requirement(
            duracao=duration if 0 < duration <= 1000 else None, **labor_requirement(order, team.executantes)).model_dump()
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid4()), "revisao": 1, "status": "rascunho", "inicio": str(team.inicio),
        "fim": str(team.fim), "setores": sorted(set(s.upper() for s in sectors)),
        "planejador": team.planejador, "atualizado_em": now, "motivo": "Criacao do rascunho",
        "integracoes": {"bma": "pendente_configuracao", "sharepoint": "pendente_configuracao"},
        "equipe": saved, "pessoas": [
            {"id": f"{m.orgid}:{m.matricula}", "nome": people[(m.orgid, m.matricula)]["nome"],
             "orgid": m.orgid, "matricula": m.matricula, "funcao": m.funcao,
             "funcao_descricao": next(c["descricao"] for c in people[(m.orgid, m.matricula)]["especialidades"]
                                      if c["codigo"] == m.funcao),
             "dias": sorted(str(day) for day in m.dias_disponiveis)}
            for m in team.executantes],
        "ordens": orders, "requisitos": requirements, "blocos": [], "indisponibilidades": [],
    }
    encoded = json.dumps(doc, ensure_ascii=False)
    pdf = schedule_pdf(doc)
    with get_connection() as conn:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute("SELECT dados_json FROM pcm_planos WHERE escopo=?", (scope,)).fetchone()
        if existing:
            return json.loads(existing["dados_json"])
        conn.execute("INSERT INTO pcm_planos VALUES (?,?,?,?,?,?)",
                     (doc["id"], scope, 1, "rascunho", now, encoded))
        conn.execute("INSERT INTO pcm_revisoes VALUES (?,?,?,?,?,?)", (doc["id"], 1, now, doc["motivo"], encoded, pdf))
    return doc


def prepare_change(current: dict, change: PlanChange) -> dict:
    if current["revisao"] != change.revisao:
        raise ConflictError("Outro usuario alterou este plano. Recarregue antes de salvar.")
    doc = copy.deepcopy(current)
    doc.update(change.model_dump(mode="json", exclude={"revisao", "confirmar", "motivo", "atualizar_especialidades"}))
    for block in doc["blocos"]:
        block["inicio"] = 0
    if set(doc["requisitos"]) != {order_key(order) for order in doc["ordens"]}:
        raise ValueError("Os requisitos devem corresponder exatamente a carteira do plano.")
    if not change.motivo.strip():
        raise ValueError("Informe o motivo da revisao.")
    if change.atualizar_especialidades:
        merge_labor(doc)
    return doc


def labor_requirement(order: dict, members) -> dict:
    rows = order.get("dados", {}).get("wplabor", [])
    if not isinstance(rows, list) or len(rows) != 1 or not isinstance(rows[0], dict):
        return {}
    row = rows[0]
    def value(key):
        raw = row.get(key)
        return raw.get("value", raw.get("_value")) if isinstance(raw, dict) else raw
    craft = str(value("craft") or "").strip()
    org = str(value("orgid") or "").strip()
    matching = {m.orgid for m in members if m.funcao == craft and (not org or m.orgid == org)}
    result = {}
    if craft and (org or len(matching) == 1):
        result.update(funcao=craft, orgid=org or next(iter(matching)))
    try:
        quantity = float(value("quantity"))
    except (TypeError, ValueError):
        quantity = 0
    if 1 <= quantity <= 100 and quantity.is_integer():
        result["pessoas"] = int(quantity)
    return result


def refresh_labor(current: dict, change: PlanChange) -> dict:
    doc = prepare_change(current, change.model_copy(update={"atualizar_especialidades": False}))
    updated = merge_labor(doc)
    return {"plano": doc, "atualizadas": updated}


def merge_labor(doc: dict) -> int:
    latest = {order_key(o): o for o in listar_ordens_trabalho(doc["setores"])}
    updated = 0
    members = PlanejamentoEquipe.model_validate(doc["equipe"]).executantes
    allocated = {b["ot"] for b in doc["blocos"]}
    for order in doc["ordens"]:
        key = order_key(order)
        fresh = latest.get(key)
        if not fresh or "wplabor" not in fresh["dados"]:
            continue
        order["dados"]["wplabor"] = copy.deepcopy(fresh["dados"]["wplabor"])
        order["sincronizado_em"] = fresh["sincronizado_em"]
        updated += 1
        req = doc["requisitos"][key]
        if key not in allocated:
            inferred = labor_requirement(order, members)
            if not req["funcao"] and not req.get("todas_especialidades"):
                req.update({k: v for k, v in inferred.items() if k in ("funcao", "orgid")})
            if not req["pessoas"] and "pessoas" in inferred:
                req["pessoas"] = inferred["pessoas"]
    if not updated:
        raise ValueError("O banco ainda nao possui mao de obra planejada para estas OTs. Atualize o banco pelo BMA primeiro.")
    return updated


def validate_plan(doc: dict, confirming: bool = False) -> None:
    people = {p["id"]: p for p in doc["pessoas"]}
    orders = {order_key(o): o for o in doc["ordens"]}
    positioned = set()
    for placement in doc.get("posicionamentos", []):
        if (placement["ot"] not in orders or placement["executante"] not in people
                or not doc["inicio"] <= placement["dia"] <= doc["fim"] or placement["ot"] in positioned):
            raise ValueError("Posicionamento de OT invalido ou duplicado.")
        positioned.add(placement["ot"])
    if confirming and positioned:
        raise ValueError("Complete os dados e as horas das OTs posicionadas no rascunho ou retire seus cartoes antes de confirmar.")
    unavailable = set()
    for item in doc["indisponibilidades"]:
        key = (item["executante"], item["dia"])
        if item["executante"] not in people or not doc["inicio"] <= item["dia"] <= doc["fim"] or key in unavailable or not item["motivo"].strip():
            raise ValueError("Indisponibilidade invalida ou duplicada.")
        unavailable.add(key)
    duration: dict[str, float] = {}
    order_weeks: dict[str, str] = {}
    ids = set()
    for block in doc["blocos"]:
        ot, day = block["ot"], block["dia"]
        if ot not in orders or block["id"] in ids:
            raise ValueError("Bloco duplicado ou OT fora da carteira.")
        ids.add(block["id"])
        req = doc["requisitos"][ot]
        week = week_start(day)
        if ot in order_weeks and order_weeks[ot] != week:
            raise ValueError(f"{orders[ot]['wonum']}: a continuacao da atividade deve permanecer na mesma semana.")
        order_weeks[ot] = week
        if not req["duracao"] or not req["pessoas"] or not specialty_selected(req):
            raise ValueError(f"{orders[ot]['wonum']}: informe duracao, pessoas e especialidade.")
        if not doc["inicio"] <= day <= doc["fim"]:
            raise ValueError("Bloco fora do horizonte.")
        earliest = day_value(text(orders[ot], "sneconstraint"))
        latest = day_value(text(orders[ot], "fnlconstraint"))
        if (earliest and day < earliest) or (latest and day > latest):
            raise ValueError(f"{orders[ot]['wonum']}: bloco viola a data de restricao do BMA.")
        members = block["executantes"]
        if len(members) != req["pessoas"] or len(set(members)) != len(members):
            raise ValueError("O bloco deve ter a quantidade simultanea de executantes requerida.")
        for member in members:
            person = people.get(member)
            if not person or not specialty_matches(person, req):
                raise ValueError("Executante fora da equipe ou sem a especialidade selecionada.")
            if day not in person["dias"] or (member, day) in unavailable:
                raise ValueError(f"{person['nome']}: indisponivel em {day}.")
        duration[ot] = duration.get(ot, 0) + block["horas"]
        if duration[ot] > req["duracao"] + 1e-6:
            raise ValueError(f"{orders[ot]['wonum']}: duracao alocada acima da atividade.")
    if confirming:
        for (person, week), total in weekly_totals([doc]).items():
            if total["horas"] > total["capacidade"] + 1e-6:
                raise ValueError(f"{people[person]['nome']}: capacidade semanal excedida na semana de {week}.")
        for key, req in doc["requisitos"].items():
            if req.get("selecionada") and (not req["duracao"] or abs(duration.get(key, 0) - req["duracao"]) > 1e-6):
                raise ValueError(f"{orders[key]['wonum']}: OT marcada para planejar ainda nao esta totalmente alocada.")
        for ot, hours in duration.items():
            if abs(hours - doc["requisitos"][ot]["duracao"]) > 1e-6:
                raise ValueError(f"{orders[ot]['wonum']}: atividade parcialmente alocada. Complete ou retire antes de confirmar.")


def save_plan(plan_id: str, change: PlanChange) -> dict:
    current = get_plan(plan_id)
    doc = prepare_change(current, change)
    if change.confirmar and current["status"] == "rascunho" and not doc["blocos"]:
        raise ValueError("Aloque ao menos uma atividade antes de confirmar.")
    validate_plan(doc, change.confirmar or current["status"] == "confirmado")
    # Recheck active workforce and open orders when confirming or revising an operational plan.
    if change.confirmar or current["status"] == "confirmado":
        calcular_equipe(PlanejamentoEquipe.model_validate(doc["equipe"]), listar_mao_de_obra(doc["setores"]))
        active = {order_key(o) for o in listar_ordens_trabalho(doc["setores"]) if o["status"].upper() not in TERMINAL}
        if any(block["ot"] not in active for block in doc["blocos"]):
            raise ValueError("Uma OT alocada foi encerrada ou saiu dos setores. Retire-a e revise o plano.")
    doc.update(revisao=current["revisao"] + 1, atualizado_em=datetime.now(timezone.utc).isoformat(),
               motivo=change.motivo.strip(),
               status="confirmado" if change.confirmar or current["status"] == "confirmado" else "rascunho")
    encoded = json.dumps(doc, ensure_ascii=False)
    pdf = schedule_pdf(doc)
    with get_connection() as conn:
        conn.execute("BEGIN IMMEDIATE")
        if doc["status"] == "confirmado":
            validate_confirmed_capacity(conn, plan_id, doc)
        updated = conn.execute(
            "UPDATE pcm_planos SET revisao=?,status=?,atualizado_em=?,dados_json=? WHERE id=? AND revisao=?",
            (doc["revisao"], doc["status"], doc["atualizado_em"], encoded, plan_id, change.revisao))
        if updated.rowcount != 1:
            raise ConflictError("O plano foi alterado. Recarregue antes de salvar.")
        conn.execute("INSERT INTO pcm_revisoes VALUES (?,?,?,?,?,?)",
                     (plan_id, doc["revisao"], doc["atualizado_em"], doc["motivo"], encoded, pdf))
    return doc


def validate_confirmed_capacity(conn, plan_id: str, doc: dict) -> None:
    rows = conn.execute("SELECT dados_json FROM pcm_planos WHERE status='confirmado' AND id<>?", (plan_id,)).fetchall()
    others = [json.loads(row["dados_json"]) for row in rows]
    for other in others:
        if (other["inicio"] <= doc["fim"] and other["fim"] >= doc["inicio"]
                and {block["ot"] for block in doc["blocos"]} & {block["ot"] for block in other["blocos"]}):
            raise ValueError("Uma OT ja esta alocada em outro planejamento confirmado no mesmo horizonte.")
    totals = weekly_totals([doc, *others])
    relevant = {(person, week_start(block["dia"])) for block in doc["blocos"] for person in block["executantes"]}
    if any(totals[key]["horas"] > totals[key]["capacidade"] + 1e-6 for key in relevant):
        raise ValueError("Executante ja comprometido: capacidade semanal excedida entre planejamentos confirmados.")


def suggest(current: dict, request: Suggestion) -> dict:
    doc = prepare_change(current, request.alteracoes)
    validate_plan(doc)
    unavailable = {(a["executante"], a["dia"]) for a in doc["indisponibilidades"]}
    # Preserve every manual/PM allocation; suggestions only fill currently free slots.
    candidates = [o for o in doc["ordens"] if mode_for(o) == request.modo]
    if request.ordens is not None:
        if any(key not in doc["requisitos"] for key in request.ordens):
            raise ValueError("Selecao de OTs fora da carteira.")
        candidates = [o for o in candidates if order_key(o) in set(request.ordens)]
    candidates.sort(key=lambda o: (doc["requisitos"][order_key(o)]["prioridade"],
                                   day_value(text(o, "targstartdate")) or "9999", o["wonum"]))
    messages = []
    day = date.fromisoformat(doc["inicio"])
    days = []
    while str(day) <= doc["fim"]:
        days.append(str(day))
        day += timedelta(days=1)
    for order in candidates:
        key = order_key(order)
        req = doc["requisitos"][key]
        if not all((req["duracao"], req["pessoas"], specialty_selected(req))):
            messages.append(f"{order['wonum']}: preparacao incompleta.")
            continue
        remaining = round(req["duracao"] - sum(b["horas"] for b in doc["blocos"] if b["ot"] == key), 6)
        activity_week = week_start(next((b["dia"] for b in doc["blocos"] if b["ot"] == key), doc["inicio"]))
        for day in days:
            if week_start(day) != activity_week:
                continue
            earliest = day_value(text(order, "sneconstraint"))
            latest = day_value(text(order, "fnlconstraint"))
            if (earliest and day < earliest) or (latest and day > latest):
                continue
            if remaining <= 1e-6:
                break
            totals = weekly_totals([doc])
            free = []
            for person in doc["pessoas"]:
                if not specialty_matches(person, req) or day not in person["dias"] or (person["id"], day) in unavailable:
                    continue
                total = totals[(person["id"], week_start(day))]
                load = sum(b["horas"] for b in doc["blocos"] if b["dia"] == day and person["id"] in b["executantes"])
                room = min(8 - load, total["capacidade"] - total["horas"])
                if room > 1e-6:
                    free.append((person["id"], room))
            if len(free) < req["pessoas"]:
                continue
            crew = free[:req["pessoas"]]
            hours = round(min(remaining, *(room for _, room in crew)), 6)
            doc["blocos"].append({"id": str(uuid4()), "ot": key, "dia": day,
                                  "inicio": 0, "horas": hours, "executantes": [p for p, _ in crew]})
            remaining = round(remaining - hours, 6)
        if remaining > 1e-6:
            messages.append(f"{order['wonum']}: faltam {remaining:g} h de duracao por falta de capacidade compativel.")
    validate_plan(doc)
    return {"blocos": doc["blocos"], "mensagens": messages}


def spreadsheet(payload: ExportOrders) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    if payload.modo == "pm" and (payload.inicio is None or payload.fim is None or payload.fim < payload.inicio):
        raise ValueError("Informe um periodo valido para exportar as preventivas.")
    available = {order_key(o): o for o in eligible(payload.setores, str(payload.inicio or ""), str(payload.fim or ""))
                 if mode_for(o) == payload.modo}
    if len(set(payload.ordens)) != len(payload.ordens) or any(key not in available for key in payload.ordens):
        raise ValueError("A exportacao contem OT fora da carteira selecionada. Recarregue os dados.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PM" if payload.modo == "pm" else "Corretiva"
    columns = [("wonum", "OT"), ("description", "Descricao"), ("status", "Status"),
               ("worktype", "Tipo"), ("maintenanceshop", "Setor"), ("location", "Local"),
               ("assetnum", "Equipamento"), ("estdur", "Duracao (h)"), ("estlabhrs", "HH estimado"),
               ("targstartdate", "Inicio alvo"), ("targcompdate", "Fim alvo"),
               ("schedstart", "Inicio programado"), ("schedfinish", "Fim programado")]
    sheet.append([label for _, label in columns])
    for key in payload.ordens:
        order = available[key]
        sheet.append([text(order, field) for field, _ in columns])
        for cell in sheet[sheet.max_row]:
            # Explicit string cells prevent spreadsheet formula execution.
            cell.data_type = "s"
        for index in range(10, 14):
            cell = sheet.cell(sheet.max_row, index)
            value = day_value(cell.value)
            if value:
                cell.value = date.fromisoformat(value)
                cell.number_format = "DD/MM/YYYY"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003B71")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate([20, 70, 18, 12, 22, 22, 22, 18, 18, 20, 20, 22, 22], 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _planner_display_name(value: str, people: list[dict]) -> str:
    planner = value.strip()
    normalized_planner = planner.casefold()
    for person in people:
        registration = str(person.get("matricula") or "").strip()
        name = str(person.get("nome") or "").strip()
        if not registration or not name:
            continue
        if normalized_planner in {registration.casefold(), name.casefold()} \
                or normalized_planner.startswith(f"{registration.casefold()} - "):
            return f"{registration} - {name}"
    return planner


def schedule_pdf(doc: dict) -> bytes:
    from services.pcm_pdf import render_schedule
    report = copy.deepcopy(doc)
    planner = str(report.get("planejador") or "").strip()
    if planner:
        people = listar_mao_de_obra(None, filtrar_nf=False)["profissionais"]
        report["planejador"] = _planner_display_name(planner, people)
    return render_schedule(report)


def get_pdf(plan_id: str, revision: int) -> bytes:
    with get_connection() as conn:
        row = conn.execute(
            "SELECT dados_json, pdf FROM pcm_revisoes WHERE plano_id=? AND revisao=?",
            (plan_id, revision),
        ).fetchone()
    if row is None:
        raise LookupError("Cronograma nao encontrado.")
    doc = json.loads(row["dados_json"])
    planner = str(doc.get("planejador") or "").strip()
    if planner:
        people = listar_mao_de_obra(None, filtrar_nf=False)["profissionais"]
        resolved = _planner_display_name(planner, people)
        if resolved != planner:
            doc["planejador"] = resolved
            from services.pcm_pdf import render_schedule
            return render_schedule(doc)
    return bytes(row["pdf"])
