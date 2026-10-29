import copy
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from models import database
from services import pcm_planning as pcm


def order(number, kind="S", **fields):
    return {"siteid": "RES", "wonum": number, "status": "PREPARED",
            "maintenanceshop": "RES-PCMTEST", "sincronizado_em": "2026-10-23T12:00:00Z",
            "dados": {"worktype": kind, "description": "PM3 teste", "targstartdate": "2026-10-26",
                      "lead": "", **fields}}


def team():
    return {"setores": ["RES-PCMTEST"], "inicio": "2026-10-26", "fim": "2026-10-27",
            "planejador": "Planejador teste", "executantes": [
                {"orgid": "BR", "matricula": key, "funcao": "180", "jornada": "adm",
                 "capacidade": "dias_8h", "dias_disponiveis": ["2026-10-26", "2026-10-27"]}
                for key in ("TEST1", "TEST2")]}


def labor():
    return {"profissionais": [{"orgid": "BR", "matricula": key, "nome": f"Pessoa {key}",
                              "setor": "RES-PCMTEST", "especialidades": [{"codigo": "180", "descricao": "Mecanica"}]}
                             for key in ("TEST1", "TEST2")]}


def prepared(duration=10, people=2, priority=3):
    return pcm.Requirement(duracao=duration, pessoas=people, orgid="BR", funcao="180",
                           liberada=True, prioridade=priority).model_dump()


def change(doc, **kwargs):
    return pcm.PlanChange(**{
        "revisao": doc["revisao"], "requisitos": doc["requisitos"], "blocos": doc["blocos"],
        "indisponibilidades": doc["indisponibilidades"], "posicionamentos": doc.get("posicionamentos", []),
        "motivo": "Teste revisao", **kwargs})


class PcmPlanningTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        db_patch = patch.object(database, "DB_PATH", Path(temp.name) / "isolated.db")
        db_patch.start()
        self.addCleanup(db_patch.stop)
        database.init_db()
        self.orders = [order("PM1"), order("E1", "E", targstartdate=""), order("OUT", targstartdate="2026-11-19")]
        for name, value in (("obter_equipe", team()), ("listar_mao_de_obra", labor()),
                            ("listar_ordens_trabalho", self.orders)):
            mock = patch.object(pcm, name, return_value=value)
            mock.start()
            self.addCleanup(mock.stop)

    def create(self):
        return pcm.create_plan(["RES-PCMTEST"])

    def test_legacy_planner_registration_is_displayed_with_name(self):
        people = [{"matricula": "B111132", "nome": "Nome do Planejador"}]
        self.assertEqual(
            pcm._planner_display_name("B111132", people),
            "B111132 - Nome do Planejador",
        )
        self.assertEqual(pcm._planner_display_name("OUTRA-MATRICULA", people), "OUTRA-MATRICULA")

    def test_create_is_idempotent_and_corrective_has_no_date_filter(self):
        doc = self.create()
        self.assertEqual([o["wonum"] for o in doc["ordens"]], ["PM1", "E1"])
        self.assertEqual(self.create()["id"], doc["id"])
        self.assertEqual(len(pcm.list_plans()), 1)
        self.assertTrue(pcm.get_pdf(doc["id"], 1).startswith(b"%PDF"))

    def test_preventive_orders_with_lead_or_scheduled_start_are_not_eligible(self):
        legacy = order("PM-LEGACY")
        del legacy["dados"]["lead"]
        self.orders.extend([
            order("PM-LEAD", lead="B114376"),
            order("PM-SCHEDULED", schedstart="2026-10-28T08:00:00-03:00"),
            legacy,
        ])
        eligible = pcm.eligible(["RES-PCMTEST"], "2026-10-26", "2026-10-27")
        self.assertEqual([item["wonum"] for item in eligible], ["PM1", "E1"])
        self.assertTrue(pcm.has_existing_planning_info(legacy))

    def test_delete_plan_removes_plan_and_all_revisions(self):
        doc = self.create()
        saved = pcm.save_plan(doc["id"], change(doc))
        self.assertEqual(len(pcm.revisions(doc["id"])), 2)
        self.assertEqual(pcm.delete_plan(doc["id"]), {"id": doc["id"], "deleted": True})
        self.assertEqual(pcm.list_plans(), [])
        with self.assertRaises(LookupError):
            pcm.get_plan(doc["id"])
        with database.get_connection() as conn:
            remaining = conn.execute("SELECT COUNT(*) FROM pcm_revisoes WHERE plano_id=?", (saved["id"],)).fetchone()[0]
        self.assertEqual(remaining, 0)

    def test_delete_missing_plan_raises_not_found(self):
        with self.assertRaisesRegex(LookupError, "Planejamento nao encontrado"):
            pcm.delete_plan("missing")

    def test_team_update_changes_only_plan_snapshot_and_creates_revision(self):
        original_team = team()
        original_team["executantes"] = original_team["executantes"][:1]
        with patch.object(pcm, "obter_equipe", return_value=original_team):
            doc = self.create()
            self.assertEqual(len(doc["pessoas"]), 1)
            updated = pcm.update_plan_team(doc["id"], pcm.PlanTeamChange(revisao=doc["revisao"], equipe=team()))
            self.assertEqual(updated["revisao"], doc["revisao"] + 1)
            self.assertEqual(len(updated["pessoas"]), 2)
            self.assertEqual(len(updated["equipe"]["executantes"]), 2)
            self.assertEqual(len(pcm.revisions(doc["id"])), 2)
            self.assertEqual(len(pcm.obter_equipe(["RES-PCMTEST"])["executantes"]), 1)

    def test_team_update_rejects_removing_an_allocated_executor(self):
        doc = self.prepared_doc()
        first_person = doc["pessoas"][0]["id"]
        block = {"id": "block-1", "ot": "RES:E1", "dia": "2026-10-26", "inicio": 0,
                 "horas": 1, "executantes": [first_person]}
        doc["blocos"] = [block]
        with database.get_connection() as conn:
            conn.execute("UPDATE pcm_planos SET dados_json=? WHERE id=?", (json.dumps(doc), doc["id"]))
        changed_team = team()
        changed_team["executantes"] = changed_team["executantes"][1:]
        with self.assertRaisesRegex(ValueError, "não pode remover executantes usados"):
            pcm.update_plan_team(doc["id"], pcm.PlanTeamChange(revisao=doc["revisao"], equipe=changed_team))

    def test_missing_requirements_are_not_invented(self):
        doc = self.create()
        self.assertIsNone(doc["requisitos"]["RES:PM1"]["duracao"])
        result = pcm.suggest(doc, pcm.Suggestion(modo="pm", alteracoes=change(doc)))
        self.assertEqual(result["blocos"], [])
        self.assertIn("preparacao", result["mensagens"][0])

    def test_positioned_card_persists_without_fabricating_hours_and_blocks_confirmation(self):
        doc = self.prepared_doc()
        doc["posicionamentos"] = [{"ot": "RES:E1", "executante": "BR:TEST1", "dia": doc["fim"]}]
        saved = pcm.save_plan(doc["id"], change(doc))
        self.assertEqual(saved["blocos"], doc["blocos"])
        self.assertEqual(pcm.get_plan(doc["id"])["posicionamentos"], doc["posicionamentos"])
        with self.assertRaisesRegex(ValueError, "posicionadas"):
            pcm.save_plan(doc["id"], change(saved, confirmar=True))
        saved["posicionamentos"] = []
        self.assertEqual(pcm.save_plan(doc["id"], change(saved, confirmar=True))["status"], "confirmado")

    def test_positioned_card_rejects_invalid_target_or_duplicate(self):
        doc = self.prepared_doc()
        for placements in (
            [{"ot": "RES:E1", "executante": "OUTSIDE", "dia": doc["fim"]}],
            [{"ot": "OUTSIDE", "executante": "BR:TEST1", "dia": doc["fim"]}],
            [{"ot": "RES:E1", "executante": "BR:TEST1", "dia": "2026-11-02"}],
            [{"ot": "RES:E1", "executante": "BR:TEST1", "dia": doc["fim"]}] * 2,
        ):
            with self.subTest(placements=placements), self.assertRaisesRegex(ValueError, "Posicionamento"):
                pcm.save_plan(doc["id"], change(doc, posicionamentos=placements))

    def test_collected_labor_populates_single_specialty_and_quantity_without_inventing_duration(self):
        self.orders[0]["dados"]["wplabor"] = [{"craft": "180", "quantity": 2, "laborhrs": 40, "orgid": "BR"}]
        doc = self.create()
        req = doc["requisitos"]["RES:PM1"]
        self.assertEqual((req["funcao"], req["orgid"], req["pessoas"]), ("180", "BR", 2))
        self.assertIsNone(req["duracao"])

    def test_refresh_labor_preserves_manual_preparation_allocations_and_persists_only_on_save(self):
        doc = self.prepared_doc()
        before = pcm.get_plan(doc["id"])
        self.orders[0]["dados"]["wplabor"] = [{"craft": "270", "quantity": 1, "laborhrs": 40, "orgid": "BR"}]
        # Existing snapshots are not silently changed.
        refreshed = pcm.refresh_labor(doc, change(doc))["plano"]
        self.assertEqual(refreshed["requisitos"]["RES:PM1"]["funcao"], "180")
        self.assertEqual(refreshed["blocos"], doc["blocos"])
        self.assertEqual(pcm.get_plan(doc["id"]), before)
        saved = pcm.save_plan(doc["id"], change(refreshed, atualizar_especialidades=True))
        self.assertEqual(saved["ordens"][0]["dados"]["wplabor"][0]["craft"], "270")
        self.assertEqual(pcm.get_plan(doc["id"])["ordens"][0]["dados"]["wplabor"], self.orders[0]["dados"]["wplabor"])

    def test_refresh_missing_collection_is_explicit_and_multiple_crafts_are_not_collapsed(self):
        doc = self.create()
        with self.assertRaisesRegex(ValueError, "Atualize o banco"):
            pcm.refresh_labor(doc, change(doc))
        self.orders[0]["dados"]["wplabor"] = [{"craft": "180", "quantity": 1}, {"craft": "270", "quantity": 2}]
        updated = pcm.refresh_labor(doc, change(doc))["plano"]
        self.assertEqual(len(updated["ordens"][0]["dados"]["wplabor"]), 2)
        self.assertEqual(updated["requisitos"]["RES:PM1"]["funcao"], "")
        self.assertIsNone(updated["requisitos"]["RES:PM1"]["pessoas"])

    def test_all_specialties_suggests_and_saves_mixed_crew_but_keeps_availability_checks(self):
        mixed_labor = labor()
        mixed_labor["profissionais"][1]["especialidades"] = [{"codigo": "210", "descricao": "Eletrica"}]
        mixed_team = team()
        mixed_team["executantes"][1]["funcao"] = "210"
        with (patch.object(pcm, "obter_equipe", return_value=mixed_team),
              patch.object(pcm, "listar_mao_de_obra", return_value=mixed_labor)):
            doc = self.create()
            doc["requisitos"]["RES:PM1"] = pcm.Requirement(
                duracao=4, pessoas=2, todas_especialidades=True, liberada=True).model_dump()
            suggested = pcm.suggest(doc, pcm.Suggestion(modo="pm", alteracoes=change(doc)))
            self.assertEqual(len(suggested["blocos"]), 1)
            self.assertEqual(set(suggested["blocos"][0]["executantes"]), {"BR:TEST1", "BR:TEST2"})
            doc["blocos"] = suggested["blocos"]
            saved = pcm.save_plan(doc["id"], change(doc, confirmar=True))
            self.assertTrue(pcm.get_plan(doc["id"])["requisitos"]["RES:PM1"]["todas_especialidades"])
            saved["indisponibilidades"] = [{"executante": "BR:TEST2", "dia": "2026-10-26", "motivo": "Falta"}]
            with self.assertRaisesRegex(ValueError, "indisponivel"):
                pcm.save_plan(doc["id"], change(saved))
            saved["indisponibilidades"] = []
            saved["requisitos"]["RES:PM1"].update(todas_especialidades=False, orgid="BR", funcao="180")
            with self.assertRaisesRegex(ValueError, "especialidade"):
                pcm.save_plan(doc["id"], change(saved))

    def test_multiday_simultaneous_crew_and_corrective_use_remaining_capacity(self):
        doc = self.create()
        doc["requisitos"]["RES:PM1"] = prepared()
        doc["requisitos"]["RES:E1"] = prepared(6, 2)
        pm = pcm.suggest(doc, pcm.Suggestion(modo="pm", alteracoes=change(doc)))
        self.assertEqual([b["horas"] for b in pm["blocos"]], [8, 2])
        self.assertTrue(all(len(b["executantes"]) == 2 for b in pm["blocos"]))
        doc["blocos"] = pm["blocos"]
        corrective = pcm.suggest(doc, pcm.Suggestion(modo="corretiva", alteracoes=change(doc)))
        self.assertEqual(corrective["blocos"][:2], pm["blocos"])
        self.assertEqual(corrective["blocos"][-1]["dia"], "2026-10-27")
        self.assertEqual(corrective["blocos"][-1]["inicio"], 0)
        self.assertEqual(corrective["blocos"][-1]["horas"], 6)
        doc["blocos"] = corrective["blocos"]
        saved = pcm.save_plan(doc["id"], change(doc, confirmar=True))
        self.assertEqual(saved["status"], "confirmado")
        self.assertEqual(saved["revisao"], 2)
        self.assertEqual(pcm.get_plan(doc["id"], 1)["blocos"], [])
        self.assertTrue(pcm.get_pdf(doc["id"], 2).startswith(b"%PDF"))
        self.assertEqual(saved["integracoes"]["bma"], "pendente_configuracao")
        with self.assertRaises(pcm.ConflictError):
            pcm.save_plan(doc["id"], change(doc))

    def prepared_doc(self):
        doc = self.create()
        doc["requisitos"]["RES:PM1"] = prepared(4, 1)
        doc["requisitos"]["RES:E1"] = prepared(4, 1)
        doc["blocos"] = [
            {"id": "a", "ot": "RES:PM1", "dia": "2026-10-26", "inicio": 0, "horas": 4,
             "executantes": ["BR:TEST1"]}]
        return doc

    def test_capacity_overlap_craft_absence_horizon_and_duplicate_crew_are_rejected(self):
        original = self.prepared_doc()
        cases = []
        doc = copy.deepcopy(original)
        doc["blocos"].append({**doc["blocos"][0], "id": "b", "ot": "RES:E1", "inicio": 3})
        # Draft overlaps are permitted; operational confirmation remains strict.
        doc = copy.deepcopy(original)
        doc["requisitos"]["RES:PM1"]["funcao"] = "OTHER"
        cases.append(doc)
        doc = copy.deepcopy(original)
        doc["indisponibilidades"] = [{"executante": "BR:TEST1", "dia": "2026-10-26", "motivo": "Falta"}]
        cases.append(doc)
        doc = copy.deepcopy(original)
        doc["blocos"][0]["dia"] = "2026-10-28"
        cases.append(doc)
        doc = copy.deepcopy(original)
        doc["blocos"][0]["executantes"] = ["BR:TEST1", "BR:TEST1"]
        cases.append(doc)
        for index, doc in enumerate(cases):
            with self.subTest(index=index), self.assertRaises(ValueError):
                pcm.save_plan(doc["id"], change(doc))
        self.assertEqual(pcm.get_plan(original["id"])["revisao"], 1)

    def test_draft_overload_persists_but_confirmation_is_blocked(self):
        doc = self.prepared_doc()
        doc["requisitos"]["RES:E1"]["duracao"] = 14
        doc["blocos"].append({**doc["blocos"][0], "id": "over", "ot": "RES:E1", "inicio": 2})
        doc["blocos"][1]["horas"] = 14
        saved = pcm.save_plan(doc["id"], change(doc))
        self.assertEqual(len(pcm.get_plan(doc["id"])["blocos"]), 2)
        with self.assertRaisesRegex(ValueError, "capacidade semanal"):
            pcm.save_plan(doc["id"], change(saved, confirmar=True))
        saved["blocos"][1]["executantes"] = ["BR:TEST2"]
        confirmed = pcm.save_plan(doc["id"], change(saved, confirmar=True))
        self.assertEqual(confirmed["status"], "confirmado")

    def test_twelve_hours_with_ten_free_saves_two_excess_only_in_draft(self):
        doc = self.prepared_doc()
        doc["requisitos"]["RES:PM1"]["duracao"] = 12
        doc["requisitos"]["RES:E1"]["duracao"] = 6
        base = doc["blocos"][0]
        doc["blocos"] = [
            {**base, "id": "existing", "ot": "RES:E1", "horas": 6},
            {**base, "id": "free2", "inicio": 6, "horas": 2},
            {**base, "id": "free8", "dia": "2026-10-27", "horas": 8},
            {**base, "id": "excess2", "horas": 2},
        ]
        saved = pcm.save_plan(doc["id"], change(doc))
        self.assertEqual(sum(b["horas"] for b in saved["blocos"]), 18)
        with self.assertRaisesRegex(ValueError, "capacidade semanal"):
            pcm.save_plan(doc["id"], change(saved, confirmar=True))

    def test_confirmed_plan_cannot_save_overloaded_operational_revision(self):
        doc = self.prepared_doc()
        saved = pcm.save_plan(doc["id"], change(doc, confirmar=True))
        saved["requisitos"]["RES:E1"]["duracao"] = 14
        saved["blocos"].append({**saved["blocos"][0], "id": "over", "ot": "RES:E1", "inicio": 2})
        saved["blocos"][1]["horas"] = 14
        with self.assertRaisesRegex(ValueError, "capacidade semanal"):
            pcm.save_plan(doc["id"], change(saved))
        persisted = pcm.get_plan(doc["id"])
        self.assertEqual(persisted["revisao"], saved["revisao"])
        self.assertEqual(persisted["status"], "confirmado")

    def test_marked_orders_must_be_allocated_and_suggestion_uses_only_selected_subset(self):
        doc = self.prepared_doc()
        doc["blocos"] = []
        doc["requisitos"]["RES:PM1"]["selecionada"] = True
        doc["requisitos"]["RES:E1"]["selecionada"] = True
        result = pcm.suggest(doc, pcm.Suggestion(modo="pm", ordens=[], alteracoes=change(doc)))
        self.assertEqual(result["blocos"], [])
        result = pcm.suggest(doc, pcm.Suggestion(modo="pm", ordens=["RES:PM1"], alteracoes=change(doc)))
        doc["blocos"] = result["blocos"]
        saved = pcm.save_plan(doc["id"], change(doc))
        self.assertTrue(saved["requisitos"]["RES:E1"]["selecionada"])
        with self.assertRaisesRegex(ValueError, "marcada"):
            pcm.save_plan(doc["id"], change(saved, confirmar=True))

    def test_daily_twelve_hours_without_start_time_can_be_confirmed(self):
        doc = self.prepared_doc()
        doc["requisitos"]["RES:PM1"]["duracao"] = 12
        doc["blocos"][0]["horas"] = 12
        doc["blocos"][0]["inicio"] = 6  # Legacy offset is discarded by new revisions.
        saved = pcm.save_plan(doc["id"], change(doc, confirmar=True))
        self.assertEqual(saved["blocos"][0]["inicio"], 0)
        self.assertEqual(saved["blocos"][0]["horas"], 12)
        self.assertEqual(saved["status"], "confirmado")

    def test_resource_checkbox_is_no_longer_required(self):
        doc = self.prepared_doc()
        doc["requisitos"]["RES:PM1"]["liberada"] = False
        pcm.validate_plan(doc, confirming=True)
        doc["blocos"] = []
        result = pcm.suggest(doc, pcm.Suggestion(modo="pm", alteracoes=change(doc)))
        self.assertEqual(sum(b["horas"] for b in result["blocos"]), 4)

    def test_capacity_is_individual_per_calendar_week_and_absences_reduce_it(self):
        from services.pcm_capacity import weekly_totals, overloaded_blocks
        doc = self.prepared_doc()
        doc["fim"] = "2026-11-02"
        doc["pessoas"][0]["dias"].append("2026-11-02")
        doc["blocos"][0]["horas"] = 18
        totals = weekly_totals([doc])
        self.assertEqual(totals[("BR:TEST1", "2026-10-26")]["capacidade"], 16)
        self.assertEqual(totals[("BR:TEST1", "2026-11-02")]["capacidade"], 8)
        self.assertEqual(overloaded_blocks(doc), {"a"})
        doc["indisponibilidades"] = [{"executante": "BR:TEST1", "dia": "2026-10-27", "motivo": "Falta"}]
        self.assertEqual(weekly_totals([doc])[("BR:TEST1", "2026-10-26")]["capacidade"], 8)

    def test_activity_cannot_continue_automatically_into_next_week(self):
        doc = self.prepared_doc()
        doc["fim"] = "2026-11-02"
        doc["pessoas"][0]["dias"].append("2026-11-02")
        doc["requisitos"]["RES:PM1"]["duracao"] = 20
        doc["blocos"] = []
        result = pcm.suggest(doc, pcm.Suggestion(modo="pm", alteracoes=change(doc)))
        self.assertEqual(sum(b["horas"] for b in result["blocos"]), 16)
        self.assertTrue(all(b["dia"] < "2026-11-02" for b in result["blocos"]))
        doc["blocos"] = result["blocos"] + [{
            **result["blocos"][0], "id": "nextweek", "dia": "2026-11-02", "horas": 4}]
        with self.assertRaisesRegex(ValueError, "mesma semana"):
            pcm.save_plan(doc["id"], change(doc))

    def test_partial_activity_can_be_draft_but_cannot_be_confirmed(self):
        doc = self.prepared_doc()
        doc["requisitos"]["RES:PM1"]["duracao"] = 6
        saved = pcm.save_plan(doc["id"], change(doc))
        with self.assertRaisesRegex(ValueError, "parcialmente"):
            pcm.save_plan(saved["id"], change(saved, confirmar=True))

    def test_absence_and_resource_postponement_create_immutable_revision(self):
        doc = self.prepared_doc()
        saved = pcm.save_plan(doc["id"], change(doc, confirmar=True))
        old_pdf = pcm.get_pdf(saved["id"], saved["revisao"])
        saved["blocos"] = [{**saved["blocos"][0], "executantes": ["BR:TEST2"]}]
        saved["indisponibilidades"] = [{"executante": "BR:TEST1", "dia": "2026-10-26", "motivo": "Falta"}]
        updated = pcm.save_plan(saved["id"], change(saved, motivo="Substituicao por falta"))
        self.assertEqual(updated["revisao"], 3)
        self.assertEqual(pcm.get_pdf(saved["id"], 2), old_pdf)
        self.assertEqual(pcm.get_plan(saved["id"], 2)["blocos"][0]["executantes"], ["BR:TEST1"])
        self.assertEqual(len(pcm.revisions(saved["id"])), 3)
        updated["blocos"] = []
        updated["requisitos"]["RES:PM1"]["liberada"] = False
        updated["requisitos"]["RES:PM1"]["observacao"] = "Recurso indisponivel; atividade postergada"
        postponed = pcm.save_plan(updated["id"], change(updated, motivo="Postergacao por recurso"))
        self.assertEqual(postponed["status"], "confirmado")
        self.assertEqual(postponed["blocos"], [])

    def test_suggestion_honors_constraint_dates_and_priority(self):
        self.orders[0]["dados"]["sneconstraint"] = "2026-10-27"
        doc = self.create()
        doc["requisitos"]["RES:PM1"] = prepared(4, 1)
        result = pcm.suggest(doc, pcm.Suggestion(modo="pm", alteracoes=change(doc)))
        self.assertEqual(result["blocos"][0]["dia"], "2026-10-27")
        doc["blocos"] = [{**result["blocos"][0], "dia": "2026-10-26"}]
        with self.assertRaisesRegex(ValueError, "restricao"):
            pcm.save_plan(doc["id"], change(doc))

    def test_overlapping_confirmed_plans_cannot_double_book_people_or_orders(self):
        doc = self.prepared_doc()
        confirmed = pcm.save_plan(doc["id"], change(doc, confirmar=True))
        other_team = team()
        other_team["fim"] = "2026-10-28"
        for member in other_team["executantes"]:
            member["dias_disponiveis"].append("2026-10-28")
        with patch.object(pcm, "obter_equipe", return_value=other_team):
            other = self.create()
        other["requisitos"]["RES:E1"] = prepared(4, 1)
        other["blocos"] = [{**confirmed["blocos"][0], "id": "other", "ot": "RES:E1"}]
        # Same day is allowed: only the combined weekly capacity is binding.
        pcm.save_plan(other["id"], change(other, confirmar=True))
        other = pcm.get_plan(other["id"])
        other["requisitos"]["RES:E1"] = prepared(22, 1)
        other["blocos"][0]["horas"] = 22
        with self.assertRaisesRegex(ValueError, "comprometido"):
            pcm.save_plan(other["id"], change(other, confirmar=True))
        other["requisitos"]["RES:PM1"] = prepared(4, 1)
        other["blocos"] = [{**confirmed["blocos"][0], "id": "other", "executantes": ["BR:TEST2"]}]
        with self.assertRaisesRegex(ValueError, "OT ja"):
            pcm.save_plan(other["id"], change(other, confirmar=True))

    def test_xlsx_is_exact_filtered_subset_and_literal_formula_with_brazilian_date(self):
        self.orders[0]["dados"]["description"] = "PM3 = literal"
        self.orders[1]["dados"]["description"] = "=HYPERLINK(\"evil\")"
        data = pcm.spreadsheet(pcm.ExportOrders(setores=["RES-PCMTEST"], inicio="2026-10-26",
                                                fim="2026-10-27", modo="corretiva", ordens=["RES:E1"]))
        sheet = load_workbook(BytesIO(data)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, "E1")
        self.assertEqual(sheet.cell(2, 2).data_type, "s")
        self.assertEqual(sheet.cell(2, 2).value, '=HYPERLINK("evil")')
        without_period = pcm.spreadsheet(pcm.ExportOrders(setores=["RES-PCMTEST"], modo="corretiva", ordens=["RES:E1"]))
        self.assertEqual(load_workbook(BytesIO(without_period)).active.max_row, 2)
        data = pcm.spreadsheet(pcm.ExportOrders(setores=["RES-PCMTEST"], inicio="2026-10-26",
                                                fim="2026-10-27", modo="pm", ordens=["RES:PM1"]))
        self.assertEqual(load_workbook(BytesIO(data)).active.cell(2, 10).number_format, "DD/MM/YYYY")
        with self.assertRaises(ValueError):
            pcm.spreadsheet(pcm.ExportOrders(setores=["RES-PCMTEST"], inicio="2026-10-26",
                                            fim="2026-10-27", modo="pm", ordens=["RES:OUT"]))

    def test_http_contract_error_mapping_and_binary_documents(self):
        from api.pcm import create, export, get, pdf, save
        from fastapi import HTTPException
        doc = create(pcm.NewPlan(setores=["RES-PCMTEST"]))
        self.assertEqual(get(doc["id"])["id"], doc["id"])
        self.assertEqual(pdf(doc["id"], 1).media_type, "application/pdf")
        self.assertEqual(export(pcm.ExportOrders(setores=["RES-PCMTEST"], inicio="2026-10-26",
                                                fim="2026-10-27", modo="pm", ordens=[])).status_code, 200)
        doc["requisitos"]["RES:PM1"] = prepared(1, 1)
        doc["blocos"] = [{"id": "a", "ot": "RES:PM1", "dia": "2026-10-26",
                          "inicio": 7, "horas": 2, "executantes": ["BR:TEST1"]}]
        with self.assertRaises(HTTPException) as invalid:
            save(doc["id"], change(doc))
        self.assertEqual(invalid.exception.status_code, 400)
        with self.assertRaises(HTTPException) as missing:
            get("unknown")
        self.assertEqual(missing.exception.status_code, 404)


if __name__ == "__main__":
    unittest.main()
