"""Opt-in integration check against the running local app; only UUID-scoped fixtures are written."""
import json
import os
import tempfile
import time
import unittest
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen
from uuid import uuid4
from zipfile import is_zipfile

from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.firefox.options import Options
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select, WebDriverWait

from models import database


@unittest.skipUnless(os.getenv("PCM_E2E") == "1", "Set PCM_E2E=1 with local servers running")
class PcmBrowserTests(unittest.TestCase):
    def api(self, path, payload=None, method=None):
        data = json.dumps(payload).encode() if payload is not None else None
        request = Request(f"http://localhost:8000/api{path}", data=data,
                          headers={"Content-Type": "application/json"}, method=method)
        with urlopen(request, timeout=30) as response:
            return json.load(response)

    def setUp(self):
        tag = uuid4().hex[:10].upper()
        self.sector = f"RES-PCMTEST-{tag}"
        self.codes = [f"PCMTEST-{tag}-{n}" for n in (1, 2)]
        self.pm = f"PCMTEST-PM-{tag}"
        self.corrective = f"PCMTEST-E-{tag}"
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.addCleanup(self.cleanup_fixtures)
        now = datetime.now(timezone.utc).isoformat()
        with database.get_connection() as conn:
            for index, code in enumerate(self.codes, 1):
                record = {"laborcode": code, "status": "ACTIVE", "laborcraftrate": [{"craft": "210", "defaultcraft": True}]}
                conn.execute(
                    "INSERT INTO mao_de_obra (orgid,laborcode,worksite,worklocation,nome,setor,dados_json,sincronizado_em) VALUES (?,?,?,?,?,?,?,?)",
                    ("BR", code, "RES", "PCMTEST", f"Executante teste {index}", self.sector, json.dumps(record), now))
            for number, kind, description, target in [
                (self.pm, "S", "PM3 atividade integrada teste", "2026-10-26"),
                (f"EXTRA-{tag}", "I", "PM2 atividade ainda nao preparada", "2026-10-26"),
                (self.corrective, "E", "Corretiva sem data alvo teste", ""),
                (f"OUT-{tag}", "S", "PM3 fora do horizonte", "2026-11-19"),
            ]:
                record = {"worktype": kind, "description": description, "targstartdate": target}
                conn.execute("INSERT INTO ordens_trabalho (siteid,wonum,status,maintenanceshop,dados_json,sincronizado_em) VALUES (?,?,?,?,?,?)",
                             ("RES", number, "PREPARED", self.sector, json.dumps(record), now))
        team = {"setores": [self.sector], "inicio": "2026-10-26", "fim": "2026-10-27", "planejador": "PCM E2E",
                "executantes": [{"orgid": "BR", "matricula": code, "funcao": "210", "jornada": "adm",
                                 "capacidade": "dias_8h", "dias_disponiveis": ["2026-10-26", "2026-10-27"]}
                                for code in self.codes]}
        self.api("/planejamento/equipe", team, "PUT")
        options = Options()
        options.add_argument("-headless")
        options.set_preference("browser.download.folderList", 2)
        options.set_preference("browser.download.dir", self.directory.name)
        options.set_preference("browser.helperApps.neverAsk.saveToDisk",
                               "application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        options.set_preference("pdfjs.disabled", True)
        self.browser = webdriver.Firefox(options=options)
        self.addCleanup(self.browser.quit)
        self.browser.set_window_size(1500, 1100)
        self.wait = WebDriverWait(self.browser, 30)

    def cleanup_fixtures(self):
        with database.get_connection() as conn:
            plans = conn.execute("SELECT id FROM pcm_planos WHERE dados_json LIKE ?", (f'%{self.sector}%',)).fetchall()
            for row in plans:
                conn.execute("DELETE FROM pcm_revisoes WHERE plano_id=?", (row["id"],))
                conn.execute("DELETE FROM pcm_planos WHERE id=?", (row["id"],))
            for table in ("planejamento_equipes", "planejamento_historico", "planejamento_padroes"):
                conn.execute(f"DELETE FROM {table} WHERE setores=?", (self.sector,))
            conn.execute("DELETE FROM ordens_trabalho WHERE maintenanceshop=?", (self.sector,))
            conn.execute("DELETE FROM mao_de_obra WHERE setor=?", (self.sector,))

    def click(self, text):
        element = self.wait.until(EC.element_to_be_clickable((By.XPATH, f"//button[contains(normalize-space(.),'{text}')]")))
        self.browser.execute_script("arguments[0].scrollIntoView({block:'center'})", element)
        element.click()

    def fill(self, label, value):
        element = self.browser.find_element(By.XPATH, f"//label[contains(normalize-space(.),'{label}')]/input")
        element.send_keys(Keys.CONTROL, "a")
        element.send_keys(str(value))
        element.send_keys(Keys.TAB)

    def downloaded(self, suffix):
        files = self.wait.until(lambda _: list(Path(self.directory.name).glob(f"*{suffix}")))
        self.wait.until(lambda _: not list(Path(self.directory.name).glob("*.part")))
        def ready(_):
            try:
                if suffix.endswith(".xlsx"):
                    return is_zipfile(files[0])
                return files[0].read_bytes().startswith(b"%PDF")
            except PermissionError:
                return False
        self.wait.until(ready)
        return files[0]

    def prepare(self, number, duration):
        self.browser.find_element(By.CSS_SELECTOR, f"button[aria-label='Editar OT {number}']").click()
        self.fill("Duração total", duration)
        craft = self.browser.find_element(By.XPATH, "//label[contains(normalize-space(.),'Especialidade exigida')]/select")
        Select(craft).select_by_value('["BR","210"]')
        for checkbox in self.browser.find_elements(By.CSS_SELECTOR, "#activity-editor input[aria-label^='Executante ']"):
            if not checkbox.is_selected():
                checkbox.click()

    def open_filters(self):
        self.browser.execute_script("document.getElementById('portfolio-filters').open = true")

    def close_editor(self):
        self.browser.find_element(By.CSS_SELECTOR, "button[aria-label='Fechar edição da atividade']").click()

    def test_collected_specialty_refresh_filter_and_real_drag_handle(self):
        plan = self.api("/pcm/planos", {"setores": [self.sector]})
        # Simulate a later successful BMA sync without touching real orders.
        with database.get_connection() as conn:
            row = conn.execute("SELECT dados_json FROM ordens_trabalho WHERE wonum=?", (self.pm,)).fetchone()
            data = json.loads(row["dados_json"])
            data.update(estdur=4, wplabor=[{"craft": "210", "orgid": "BR", "quantity": 1, "laborhrs": 4}])
            conn.execute("UPDATE ordens_trabalho SET dados_json=? WHERE wonum=?", (json.dumps(data), self.pm))
        self.browser.set_window_size(1600, 1900)
        self.browser.get(f"http://localhost:3000/planejamento/alocacao?plano={plan['id']}")
        self.click("Reler especialidades do banco")
        self.wait.until(lambda d: "relidas do banco local" in d.find_element(By.TAG_NAME, "body").text)
        # Duration is intentionally not replaced when refreshing specialty metadata.
        self.browser.find_element(By.CSS_SELECTOR, f"button[aria-label='Editar OT {self.pm}']").click()
        self.fill("Duração total", 4)
        self.close_editor()
        self.open_filters()
        specialty_filter = self.browser.find_element(By.CSS_SELECTOR, "select[aria-label='Filtrar carteira por especialidade']")
        Select(specialty_filter).select_by_value("210")
        self.assertTrue(self.browser.find_elements(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']"))
        self.assertIn("Especialidade BMA:", self.browser.find_element(By.TAG_NAME, "body").text)
        source = self.browser.find_element(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']")
        target = self.browser.find_element(By.CSS_SELECTOR, "th[data-person]")
        ActionChains(self.browser).move_to_element(source).click_and_hold().pause(.5).move_to_element(target).pause(.5).release().perform()
        self.wait.until(lambda d: bool(d.find_elements(By.XPATH, f"//tbody/tr[1]/td[1]//button[contains(.,'{self.pm}')]")))
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']"))
        self.browser.find_element(By.XPATH, f"//td//button[contains(.,'{self.pm}')]").click()
        self.click("Retirar OT para reencaixe")
        self.wait.until(EC.alert_is_present()).accept()
        self.close_editor()
        self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']")))
        source = self.browser.find_element(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']")
        target = self.browser.find_element(By.CSS_SELECTOR, "td[data-day='2026-10-27']")
        ActionChains(self.browser).move_to_element(source).click_and_hold().pause(.5).move_to_element(target).pause(.5).release().perform()
        self.wait.until(lambda d: bool(d.find_elements(By.XPATH, f"//tbody/tr[1]/td[2]//button[contains(.,'{self.pm}')]")))
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']"))
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 2" in d.find_element(By.TAG_NAME, "header").text)
        saved = self.api(f"/pcm/planos/{plan['id']}")
        self.assertEqual(saved["blocos"][0]["dia"], "2026-10-27")
        self.assertEqual(saved["blocos"][0]["executantes"], [f"BR:{self.codes[0]}"])
        self.assertEqual(next(o for o in saved["ordens"] if o["wonum"] == self.pm)["dados"]["wplabor"][0]["craft"], "210")
        self.browser.get("http://localhost:3000/planejamento")
        self.wait.until(EC.presence_of_element_located((By.XPATH, "//button[contains(.,'banco de dados')]")))

    def test_missing_preparation_stays_visible_in_exact_drop_cell_until_completed(self):
        plan = self.api("/pcm/planos", {"setores": [self.sector]})
        self.browser.set_window_size(1600, 1900)
        self.browser.get(f"http://localhost:3000/planejamento/alocacao?plano={plan['id']}")
        source = self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']")))
        self.assertFalse(self.browser.find_elements(By.ID, "activity-editor"))
        self.browser.find_element(By.CSS_SELECTOR, "input[aria-label='Selecionar todas as OTs filtradas']").click()
        self.browser.find_element(By.CSS_SELECTOR, "input[aria-label='Selecionar todas as OTs filtradas']").click()
        self.browser.find_element(By.CSS_SELECTOR, f"input[aria-label='Planejar {self.pm}']").click()
        self.assertIn("1 OT(s) pendentes marcadas", self.browser.find_element(By.TAG_NAME, "body").text)
        target = self.browser.find_element(By.CSS_SELECTOR, "td[data-day='2026-10-27']")
        ActionChains(self.browser).move_to_element(source).click_and_hold().pause(.5).move_to_element(target).pause(.5).release().perform()
        self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, f"td[data-day='2026-10-27'] [data-placement-ot='RES:{self.pm}']")))
        card = self.browser.find_element(By.CSS_SELECTOR, f"[data-placement-ot='RES:{self.pm}']")
        self.assertEqual(card.text, f"{self.pm} - PM\nPM3 atividade integrada teste")
        meter = self.browser.find_element(By.CSS_SELECTOR, "tbody tr:first-child td[data-day='2026-10-27'] [role=progressbar]")
        self.assertEqual(meter.get_attribute("aria-valuenow"), "0")
        self.assertEqual(meter.get_attribute("aria-valuetext"), "0% · 0 h ocupadas · 8 h disponíveis")
        if os.getenv("PCM_COMPACT_SCREENSHOT"):
            self.browser.execute_script("arguments[0].scrollIntoView({block:'center'})", card)
            self.browser.save_screenshot(os.environ["PCM_COMPACT_SCREENSHOT"])
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']"))
        self.assertIn("0 OT(s) pendentes marcadas", self.browser.find_element(By.TAG_NAME, "body").text)
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 2" in d.find_element(By.TAG_NAME, "header").text)
        saved = self.api(f"/pcm/planos/{plan['id']}")
        self.assertEqual(saved["blocos"], [])
        self.assertTrue(saved["requisitos"][f"RES:{self.pm}"]["selecionada"])
        self.assertEqual(saved["posicionamentos"], [{"ot": f"RES:{self.pm}", "executante": f"BR:{self.codes[0]}", "dia": "2026-10-27"}])
        self.browser.refresh()
        self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, f"[data-placement-ot='RES:{self.pm}']")))
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']"))
        self.browser.find_element(By.CSS_SELECTOR, f"[data-placement-ot='RES:{self.pm}'] button[draggable]").click()
        self.fill("Duração total", 4)
        Select(self.browser.find_element(By.XPATH, "//label[contains(normalize-space(.),'Especialidade exigida')]/select")).select_by_value('["BR","210"]')
        for checkbox in self.browser.find_elements(By.CSS_SELECTOR, "#activity-editor input[aria-label^='Executante ']"):
            if not checkbox.is_selected():
                checkbox.click()
        self.assertFalse(self.browser.find_elements(By.XPATH, "//label[contains(.,'Pessoas simultâneas')]/input"))
        self.assertIn("2 selecionada(s)", self.browser.find_element(By.CSS_SELECTOR, "output[aria-label='Pessoas simultâneas']").text)
        self.click("Adicionar bloco")
        self.assertFalse(self.browser.find_elements(By.ID, "activity-editor"))
        self.wait.until(lambda d: bool(d.find_elements(By.XPATH, f"//tbody/tr[1]/td[2]//button[contains(.,'{self.pm}')]")))
        meter = self.browser.find_element(By.CSS_SELECTOR, "tbody tr:first-child td[data-day='2026-10-27'] [role=progressbar]")
        self.assertEqual(meter.get_attribute("aria-valuenow"), "50")
        self.assertEqual(meter.get_attribute("aria-valuetext"), "50% · 4 h ocupadas · 4 h disponíveis")
        allocated = self.browser.find_element(By.XPATH, f"//tbody/tr[1]/td[2]//button[contains(.,'{self.pm}')]")
        self.assertEqual(allocated.text, f"{self.pm} - PM\nPM3 atividade integrada teste")
        self.assertFalse(self.browser.find_elements(By.XPATH, f"//tbody/tr[1]/td[1]//button[contains(.,'{self.pm}')]"))
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, f"[data-drag-ot='RES:{self.pm}']"))
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 3" in d.find_element(By.TAG_NAME, "header").text)
        saved = self.api(f"/pcm/planos/{plan['id']}")
        self.assertEqual(saved["blocos"][0]["dia"], "2026-10-27")
        self.assertEqual(saved["posicionamentos"], [])

    def test_nfe_report_button_waits_for_sync_and_never_downloads_on_failure(self):
        self.browser.get(f"http://localhost:3000/planejamento/ots?setor={self.sector}")
        self.wait.until(EC.element_to_be_clickable((By.XPATH, "//button[contains(.,'Atualizar BMA e gerar Excel NFE')]")))
        self.browser.execute_script("""
            const original = window.fetch.bind(window);
            window.nfeCalls = { polls: 0, exports: 0 };
            window.fetch = async (input, init) => {
                const url = String(input);
                if (url.endsWith('/api/planejamento/sincronizar')) {
                    window.nfeCalls.payload = JSON.parse(init.body);
                    return new Response(JSON.stringify({ job_id: 'NFE-BROWSER-TEST' }), { status: 202 });
                }
                if (url.endsWith('/api/planejamento/sincronizar/NFE-BROWSER-TEST')) {
                    window.nfeCalls.polls++;
                    const saving = window.nfeCalls.polls === 1;
                    return new Response(JSON.stringify({
                        status: saving ? 'saving' : 'failed', ots_lidas: 2,
                        mensagem: saving ? 'Gravando teste' : 'Falha de coleta teste',
                        erro: saving ? null : 'Falha de coleta teste'
                    }), { status: 200 });
                }
                if (url.includes('/api/pcm/exportar-nfe/')) window.nfeCalls.exports++;
                return original(input, init);
            };
        """)
        self.click("Atualizar BMA e gerar Excel NFE")
        self.wait.until(lambda d: "Gravando teste" in d.find_element(By.TAG_NAME, "body").text)
        button = self.browser.find_element(By.XPATH, "//button[contains(.,'Atualizando BMA')]")
        self.assertFalse(button.is_enabled())
        self.wait.until(lambda d: any("Falha de coleta teste" in e.text for e in d.find_elements(By.CSS_SELECTOR, "[role=alert]")))
        calls = self.browser.execute_script("return window.nfeCalls")
        self.assertEqual(calls["payload"], {"setores": ["RES-TC-NF"]})
        self.assertEqual(calls["polls"], 2)
        self.assertEqual(calls["exports"], 0)
        self.assertFalse(list(Path(self.directory.name).glob("*.xlsx")))

    def test_pencil_editor_uses_bma_defaults_dates_and_selected_crew_with_daily_continuation(self):
        with database.get_connection() as conn:
            for number, duration in ((self.pm, 8), (self.corrective, 6.4)):
                row = conn.execute("SELECT dados_json FROM ordens_trabalho WHERE wonum=?", (number,)).fetchone()
                data = json.loads(row["dados_json"])
                data.update(estdur=duration, targstartdate="2026-10-26", targcompdate="2026-10-27",
                            wplabor=[{"craft": "210", "orgid": "BR", "quantity": 1, "laborhrs": duration}])
                conn.execute("UPDATE ordens_trabalho SET dados_json=? WHERE wonum=?", (json.dumps(data), number))
        plan = self.api("/pcm/planos", {"setores": [self.sector]})
        self.assertEqual(plan["requisitos"][f"RES:{self.pm}"]["pessoas"], 1)
        payload = {key: plan[key] for key in ("revisao", "requisitos", "blocos", "indisponibilidades")}
        payload.update(motivo="Fixture de carga parcial", blocos=[{
            "id": uuid4().hex, "ot": f"RES:{self.corrective}", "dia": "2026-10-26",
            "inicio": 0, "horas": 6.4, "executantes": [f"BR:{self.codes[0]}"],
        }])
        self.api(f"/pcm/planos/{plan['id']}", payload, "PUT")
        self.browser.get(f"http://localhost:3000/planejamento/alocacao?plano={plan['id']}")
        pencil = self.wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, f"button[aria-label='Editar OT {self.pm}']")))
        self.assertFalse(self.browser.find_elements(By.ID, "activity-editor"))
        pencil.click()
        duration = self.browser.find_element(By.XPATH, "//label[contains(.,'Duração total')]/input")
        self.assertEqual(float(duration.get_attribute("value")), 8)
        self.assertIn("1 selecionada(s)", self.browser.find_element(By.CSS_SELECTOR, "output[aria-label='Pessoas simultâneas']").text)
        self.browser.find_element(By.CSS_SELECTOR, "input[aria-label='Executante Executante teste 2']").click()
        self.assertIn("2 selecionada(s)", self.browser.find_element(By.CSS_SELECTOR, "output[aria-label='Pessoas simultâneas']").text)
        for day in ("26/10/2026", "27/10/2026"):
            checkbox = self.browser.find_element(By.CSS_SELECTOR, f"input[aria-label='Executar em {day}']")
            self.assertTrue(checkbox.is_selected())
            self.assertIn("text-emerald-200", checkbox.find_element(By.XPATH, "..").get_attribute("class"))
        self.click("Adicionar bloco")
        self.assertFalse(self.browser.find_elements(By.ID, "activity-editor"))
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 3" in d.find_element(By.TAG_NAME, "header").text)
        saved = self.api(f"/pcm/planos/{plan['id']}")
        blocks = [b for b in saved["blocos"] if b["ot"] == f"RES:{self.pm}"]
        self.assertEqual([(b["dia"], b["horas"]) for b in blocks], [("2026-10-26", 1.6), ("2026-10-27", 6.4)])
        self.assertTrue(all(b["executantes"] == [f"BR:{code}" for code in self.codes] for b in blocks))
        self.assertEqual(saved["requisitos"][f"RES:{self.pm}"]["pessoas"], 2)

    def test_gallery_filters_batch_drag_and_reschedule_sequence_persist_after_reload(self):
        with database.get_connection() as conn:
            rows = conn.execute("SELECT wonum,dados_json FROM ordens_trabalho WHERE maintenanceshop=?", (self.sector,)).fetchall()
            pm_numbers = []
            for row in rows:
                data = json.loads(row["dados_json"])
                if data["worktype"] in ("I", "S") and not row["wonum"].startswith("OUT-"):
                    pm_numbers.append(row["wonum"])
                    data.update(estdur=6, wplabor=[{"craft": "210", "orgid": "BR", "quantity": 1, "laborhrs": 6}])
                    conn.execute("UPDATE ordens_trabalho SET dados_json=? WHERE wonum=?", (json.dumps(data), row["wonum"]))
        plan = self.api("/pcm/planos", {"setores": [self.sector]})
        self.browser.set_window_size(1600, 1050)
        self.browser.get(f"http://localhost:3000/planejamento/alocacao?plano={plan['id']}")
        self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "section[aria-label='Galeria vertical de OTs']")))
        self.assertFalse(self.browser.find_elements(By.ID, "activity-editor"))
        aligned = self.browser.execute_script("""
            const gallery = document.querySelector('[aria-label="Galeria vertical de OTs"]').getBoundingClientRect();
            const grid = document.getElementById('integrated-schedule').getBoundingClientRect();
            return Math.abs(gallery.top-grid.top)<10 && gallery.right<grid.left;
        """)
        self.assertTrue(aligned)
        self.open_filters()
        category = Select(self.browser.find_element(By.CSS_SELECTOR, "select[aria-label='Filtrar tipo de OT']"))
        specialty = Select(self.browser.find_element(By.CSS_SELECTOR, "select[aria-label='Filtrar carteira por especialidade']"))
        category.select_by_value("pm3")
        specialty.select_by_value("210")
        self.assertEqual(len(self.browser.find_elements(By.CSS_SELECTOR, "[data-drag-ot]")), 1)
        self.browser.find_element(By.CSS_SELECTOR, "input[aria-label='Selecionar todas as OTs filtradas']").click()
        category.select_by_value("pm2")
        self.assertEqual(len(self.browser.find_elements(By.CSS_SELECTOR, "[data-drag-ot]")), 1)
        self.assertIn("0 OT(s) pendentes marcadas", self.browser.find_element(By.TAG_NAME, "body").text)
        self.browser.find_element(By.CSS_SELECTOR, "input[aria-label='Selecionar todas as OTs filtradas']").click()
        category.select_by_value("")
        self.assertIn("2 OT(s) pendentes marcadas", self.browser.find_element(By.TAG_NAME, "body").text)
        self.browser.execute_script("document.getElementById('portfolio-filters').open = false")
        source = self.browser.find_element(By.CSS_SELECTOR, "[data-drag-ot]")
        target = self.browser.find_element(By.CSS_SELECTOR, "th[data-person]")
        self.browser.execute_script("arguments[0].scrollIntoView({block:'center'})", source)
        ActionChains(self.browser).move_to_element(source).click_and_hold().pause(.4).move_to_element(target).pause(.4).release().perform()
        self.wait.until(lambda d: len(d.find_elements(By.CSS_SELECTOR, "tbody tr:first-child [data-schedule-block]")) == 3)
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, "[data-drag-ot]"))
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 2" in d.find_element(By.TAG_NAME, "header").text)
        saved = self.api(f"/pcm/planos/{plan['id']}")
        first = saved["blocos"][0]
        last = saved["blocos"][-1]
        self.assertNotEqual(first["ot"], last["ot"])
        source = self.browser.find_element(By.CSS_SELECTOR, f"tbody tr:first-child [data-schedule-block='{last['id']}']")
        target = self.browser.find_element(By.CSS_SELECTOR, f"tbody tr:first-child [data-schedule-block='{first['id']}']")
        ActionChains(self.browser).move_to_element(source).click_and_hold().pause(.4).move_to_element(target).pause(.4).release().perform()
        self.wait.until(lambda d: "Sequência reprogramada" in d.find_element(By.TAG_NAME, "body").text)
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 3" in d.find_element(By.TAG_NAME, "header").text)
        updated = self.api(f"/pcm/planos/{plan['id']}")
        self.assertEqual([(b["ot"], b["dia"], b["horas"]) for b in updated["blocos"]], [
            (last["ot"], "2026-10-26", 6), (first["ot"], "2026-10-26", 2), (first["ot"], "2026-10-27", 4),
        ])
        self.assertEqual(sum(b["horas"] for b in updated["blocos"]), 12)
        self.browser.refresh()
        first_card = self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "tbody tr:first-child [data-schedule-block]")))
        self.assertIn(last["ot"].split(":")[1], first_card.text)
        if os.getenv("PCM_GALLERY_SCREENSHOT"):
            self.browser.execute_script("document.getElementById('integrated-schedule').scrollIntoView({block:'start'})")
            self.browser.save_screenshot(os.environ["PCM_GALLERY_SCREENSHOT"])

    def test_integrated_planning_export_confirm_absence_and_history(self):
        self.browser.get(f"http://localhost:3000/planejamento/ots?setor={self.sector}")
        self.wait.until(EC.element_to_be_clickable((By.XPATH, "//button[normalize-space(.)='Planejar PM']")))
        number_filter = self.browser.find_element(By.CSS_SELECTOR, "input[aria-label='Filtrar Ordem de Trabalho']")
        number_filter.send_keys(self.pm)
        self.click("Exportar Excel filtrado")
        sheet = load_workbook(self.downloaded(".xlsx")).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, self.pm)
        number_filter.send_keys(Keys.CONTROL, "a")
        number_filter.send_keys(Keys.BACKSPACE)
        corrective_export = self.browser.find_element(By.XPATH, "//article[.//h2[normalize-space(.)='Planejamento Corretivo']]//button[contains(normalize-space(.),'Exportar Excel')]")
        self.browser.execute_script("arguments[0].scrollIntoView({block:'center'})", corrective_export)
        corrective_export.click()
        sheet = load_workbook(self.downloaded("corretiva-2026-10-26.xlsx")).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, self.corrective)
        self.click("Planejar PM")
        self.wait.until(lambda d: "/alocacao?" in d.current_url)
        self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "input[aria-label='Buscar atividade']")))
        for removed in ("Conferi recursos", "Restrições / motivo", "Acompanhamento · faltas",
                        "Motivo / descrição desta revisão", "Histórico e cronogramas PDF"):
            self.assertNotIn(removed, self.browser.find_element(By.TAG_NAME, "body").text)
        self.prepare(self.pm, 10)
        craft = self.browser.find_element(By.XPATH, "//label[contains(normalize-space(.),'Especialidade exigida')]/select")
        Select(craft).select_by_value("todos")
        for checkbox in self.browser.find_elements(By.CSS_SELECTOR, "#activity-editor input[aria-label^='Executante ']"):
            if not checkbox.is_selected():
                checkbox.click()
        self.assertEqual(len(self.browser.find_elements(By.XPATH, "//label[contains(normalize-space(.),'Executante teste')]/input")), 2)
        self.close_editor()
        self.browser.find_element(By.CSS_SELECTOR, f"input[aria-label='Planejar {self.pm}']").click()
        self.open_filters()
        specialty_filter = self.browser.find_element(By.CSS_SELECTOR, "select[aria-label='Filtrar carteira por especialidade']")
        Select(specialty_filter).select_by_value("indefinida")
        self.assertFalse(self.browser.find_elements(By.XPATH, f"//button[@draggable='true'][contains(.,'{self.pm}')]"))
        Select(specialty_filter).select_by_value("qualquer")
        source = self.browser.find_element(By.XPATH, f"//button[@draggable='true'][contains(.,'{self.pm}')]")
        target = self.browser.find_element(By.CSS_SELECTOR, "section[aria-label='Grade de capacidade por integrante e dia'] tbody tr td")
        self.browser.execute_script("""
            const data = new DataTransfer();
            arguments[0].dispatchEvent(new DragEvent('dragstart', {dataTransfer:data, bubbles:true}));
            arguments[1].dispatchEvent(new DragEvent('dragover', {dataTransfer:data, bubbles:true, cancelable:true}));
            arguments[1].dispatchEvent(new DragEvent('drop', {dataTransfer:data, bubbles:true, cancelable:true}));
        """, source, target)
        self.wait.until(lambda d: "posicionada(s)" in d.find_element(By.TAG_NAME, "body").text)
        self.assertTrue(self.browser.find_elements(By.XPATH, f"//tbody/tr[1]/td[1]//button[contains(.,'{self.pm}')]"))
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, f"input[aria-label='Planejar {self.pm}']"))
        Select(specialty_filter).select_by_value("")
        self.assertFalse(self.browser.find_elements(By.XPATH, "//label[contains(.,'Início relativo')]"))
        pm_blocks = self.browser.find_elements(By.XPATH, f"//tbody/tr[1]/td//button[contains(.,'{self.pm}')]")
        pm_blocks[-1].click()
        Select(self.browser.find_element(By.XPATH, "//label[normalize-space(text())='Dia']/select")).select_by_value("2026-10-26")
        self.click("Aplicar edição do bloco")
        meter = self.browser.find_element(By.CSS_SELECTOR, "tbody tr:first-child td[data-day='2026-10-26'] [role=progressbar]")
        self.assertEqual(meter.get_attribute("aria-valuenow"), "100")
        self.assertEqual(meter.get_attribute("aria-valuetext"), "125% · 10 h ocupadas · 0 h disponíveis")
        self.assertNotIn("capacidade semanal excedida.", self.browser.find_element(By.TAG_NAME, "body").text)
        self.browser.find_element(By.LINK_TEXT, "Acompanhamento").click()
        self.wait.until(EC.alert_is_present()).dismiss()
        self.assertIn("/alocacao?", self.browser.current_url)
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 2" in d.find_element(By.TAG_NAME, "header").text)
        self.browser.refresh()
        self.wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, "input[aria-label='Buscar atividade']")))
        self.assertFalse(self.browser.find_elements(By.CSS_SELECTOR, f"input[aria-label='Planejar {self.pm}']"))
        self.click("Encaixar corretiva")
        self.prepare(self.corrective, 8)
        self.close_editor()
        self.browser.find_element(By.CSS_SELECTOR, f"input[aria-label='Planejar {self.corrective}']").click()
        self.browser.execute_script("document.querySelector('select[aria-label=\"Executante para distribuir lote\"]').closest('details').open = true")
        Select(self.browser.find_element(By.CSS_SELECTOR, "select[aria-label='Executante para distribuir lote']")).select_by_index(1)
        self.click("Distribuir marcadas no período")
        self.wait.until(lambda d: "2 h pendentes" in d.find_element(By.TAG_NAME, "body").text)
        self.browser.find_element(By.XPATH, f"//td//button[contains(.,'{self.corrective}')]").click()
        self.fill("Horas da atividade neste dia", 8)
        self.click("Aplicar edição do bloco")
        self.wait.until(lambda d: "alocação(ões) com capacidade semanal excedida" in d.find_element(By.TAG_NAME, "body").text)
        self.click("Salvar rascunho")
        self.wait.until(lambda d: "revisão 3" in d.find_element(By.TAG_NAME, "header").text)
        self.browser.refresh()
        self.wait.until(lambda d: "alocação(ões) com capacidade semanal excedida" in d.find_element(By.TAG_NAME, "body").text)
        self.click("Revisar e confirmar")
        self.browser.find_element(By.XPATH, "//label[contains(normalize-space(.),'Conferi atividades incluídas')]/input").click()
        self.assertFalse(self.browser.find_element(By.XPATH, "//button[contains(.,'Confirmar planejamento local')]").is_enabled())
        self.click("Encaixar corretiva")
        self.browser.find_element(By.XPATH, f"//td//button[contains(.,'{self.corrective}')]").click()
        self.click("Retirar OT para reencaixe")
        self.wait.until(EC.alert_is_present()).accept()
        self.fill("Duração total", 6)
        self.close_editor()
        self.click("Sugerir encaixe")
        self.wait.until(lambda d: "Sugestão aplicada" in d.find_element(By.TAG_NAME, "body").text)
        self.click("Revisar e confirmar")
        self.browser.find_element(By.XPATH, "//label[contains(normalize-space(.),'Conferi atividades incluídas')]/input").click()
        self.click("Confirmar planejamento local")
        self.wait.until(lambda d: "confirmado · revisão 4" in d.find_element(By.TAG_NAME, "header").text)
        plan_id = self.browser.current_url.split("plano=")[1].split("&")[0]
        saved = self.api(f"/pcm/planos/{plan_id}")
        self.assertTrue(saved["requisitos"][f"RES:{self.pm}"]["todas_especialidades"])
        self.assertEqual(len(saved["blocos"]), 3)
        self.assertEqual(sum(b["horas"] * len(b["executantes"]) for b in saved["blocos"]), 32)
        self.assertEqual(saved["motivo"], "Atualização do planejamento")
        if os.getenv("PCM_SCREENSHOT"):
            self.browser.execute_script("window.scrollTo(0,0)")
            self.browser.save_screenshot(os.environ["PCM_SCREENSHOT"])
        self.click("PDF da revisão salva")
        self.assertTrue(self.downloaded(".pdf").read_bytes().startswith(b"%PDF"))
        # Open the remaining allocation rather than the activity card.
        grid_block = self.browser.find_element(By.XPATH, f"//td//button[contains(normalize-space(.),'{self.corrective}')]")
        grid_block.click()
        self.fill("Horas da atividade neste dia", 6)
        self.click("Aplicar edição do bloco")
        self.click("Salvar reprogramação")
        self.wait.until(lambda d: "revisão 5" in d.find_element(By.TAG_NAME, "header").text)
        updated = self.api(f"/pcm/planos/{plan_id}")
        self.assertEqual(len(updated["blocos"]), 3)
        self.assertTrue(all(b["inicio"] == 0 for b in updated["blocos"]))
        self.assertEqual(len(self.api(f"/pcm/planos/{plan_id}/revisoes")), 5)
        for width in (768, 375, 1920):
            self.browser.set_window_size(width, 1024)
            time.sleep(.2)
            self.assertFalse(self.browser.execute_script("return document.documentElement.scrollWidth > innerWidth + 1"), f"Horizontal overflow at {width}px")
        self.browser.get("http://localhost:3000/planejamento/acompanhamento")
        self.wait.until(lambda d: self.sector in d.find_element(By.TAG_NAME, "body").text)
        self.assertIn("Confirmado localmente", self.browser.find_element(By.TAG_NAME, "body").text)
        self.browser.set_window_size(768, 1024)
        time.sleep(.3)
        self.assertFalse(self.browser.execute_script("return document.documentElement.scrollWidth > innerWidth + 1"))
