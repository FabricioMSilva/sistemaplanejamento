import unittest
from datetime import date, time
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from services import nfe_report as report
from services.pcm_planning import ConflictError


def order(**fields):
    return {
        "siteid": "RES", "wonum": "RES001", "status": "INPRG",
        "maintenanceshop": report.NF_SECTOR, "sincronizado_em": "2026-10-23T14:01:00+00:00",
        "dados": {
            "description": "=literal", "worktype": "A", "location": "LOCAL",
            "origin": "3 AVAIL", "planner": "RES-TC-NF-REF-PM",
            "reportdate": "2026-10-23T01:30:00Z", "sneconstraint": "2026-10-26T08:00:00-03:00",
            "fnlconstraint": "2026-10-30T16:00:00-03:00", "pmduedate": "2026-10-28",
            "schedstart": "2026-10-26T08:00:00-03:00", "schedfinish": "2026-10-27T12:00:00-03:00",
            "estlabhrs": 40, "estdur": 0, "pmnum": "PM001", "jpnum": "JP001",
            "description_longdescription": "<p>Primeira &amp; segunda</p><div>Outra<br>linha</div>",
            **fields,
        },
    }


class NfeReportTests(unittest.TestCase):
    def workbook(self, orders):
        return load_workbook(BytesIO(report.render_report(orders)))

    def test_exact_physical_columns_merges_headers_and_field_mapping(self):
        workbook = self.workbook([order()])
        self.assertEqual(workbook.sheetnames, ["Report"])
        sheet = workbook["Report"]
        self.assertEqual(sheet.max_column, 29)
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["A2"].value, "Planejamento NFE")
        self.assertEqual({c.coordinate[:-1]: c.value for c in sheet[3] if c.value}, report.HEADERS)
        self.assertEqual({str(m) for m in sheet.merged_cells}, {
            "A1:G1", "H1:P1", "A2:AC2", "B3:C3", "G3:H3", "K3:L3", "M3:Q3", "R3:S3",
            "W3:X3", "Y3:Z3", "B4:C4", "G4:H4", "M4:O4", "P4:Q4",
        })
        for address, expected in {
            "A4": "RES001", "B4": "=literal", "D4": "LOCAL", "E4": "INPRG", "F4": "A",
            "G4": "3 AVAIL", "I4": report.NF_SECTOR, "J4": "RES-TC-NF-REF-PM",
            "AA4": "PM001", "AB4": "JP001",
        }.items():
            self.assertEqual(sheet[address].value, expected)
            self.assertEqual(sheet[address].data_type, "s")
        self.assertEqual(sheet["K4"].value.date(), date(2026, 10, 22))
        self.assertEqual(sheet["L4"].value, time(22, 30))
        self.assertEqual(sheet["M4"].value.date(), date(2026, 10, 26))
        self.assertEqual(sheet["P4"].value, time(8))
        self.assertEqual(sheet["R4"].value.date(), date(2026, 10, 30))
        self.assertEqual(sheet["S4"].value, time(16))
        self.assertEqual(sheet["T4"].value.date(), date(2026, 10, 28))
        self.assertEqual(sheet["W4"].value.date(), date(2026, 10, 26))
        self.assertEqual(sheet["X4"].value, time(8))
        self.assertEqual(sheet["Y4"].value.date(), date(2026, 10, 27))
        self.assertEqual(sheet["Z4"].value, time(12))
        self.assertAlmostEqual(sheet["U4"].value.total_seconds() / 3600, 40)
        self.assertEqual(sheet["U4"].number_format, "[h]:mm")
        self.assertEqual(sheet["V4"].value.total_seconds(), 0)
        self.assertEqual(sheet["AC4"].value, "Primeira & segunda\n\nOutra\nlinha")
        self.assertEqual(sheet.freeze_panes, "A4")
        self.assertEqual(sheet.auto_filter.ref, "A3:AC4")

    def test_missing_fields_remain_empty_without_invented_values(self):
        sheet = self.workbook([{"wonum": "RES002", "dados": {}}])["Report"]
        for col in ("G", "J", "K", "L", "U", "V", "W", "X"):
            self.assertIsNone(sheet[f"{col}4"].value)
        self.assertEqual(self.workbook([])["Report"].max_row, 3)

    def test_invalid_dates_and_hours_fail_explicitly(self):
        for fields in ({"reportdate": "invalid"}, {"estdur": "invalid"}, {"estdur": "nan"},
                       {"estlabhrs": -1}, {"estlabhrs": "inf"}):
            with self.subTest(fields=fields), self.assertRaises(ValueError):
                report.render_report([order(**fields)])

    @patch.object(report, "listar_ordens_trabalho")
    @patch.object(report, "obter_sincronizacao")
    def test_export_only_includes_rows_from_successful_nf_sync(self, sync, rows):
        sync.return_value = {
            "tipo": "ots", "setores": [report.NF_SECTOR], "status": "completed", "ots_salvas": 1,
            "iniciada_em": "2026-10-23T14:00:00+00:00", "concluida_em": "2026-10-23T14:02:00+00:00",
        }
        old = {**order(), "wonum": "RES-OLD", "sincronizado_em": "2026-10-22T14:01:00+00:00"}
        rows.return_value = [old, order()]
        sheet = load_workbook(BytesIO(report.export_synced_report("job")))["Report"]
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["A4"].value, "RES001")
        rows.assert_called_once_with([report.NF_SECTOR])
        rows.return_value = [old]
        with self.assertRaises(ConflictError):
            report.export_synced_report("job")
        sync.return_value["ots_salvas"] = 0
        self.assertEqual(load_workbook(BytesIO(report.export_synced_report("job")))["Report"].max_row, 3)

    @patch.object(report, "listar_ordens_trabalho")
    @patch.object(report, "obter_sincronizacao")
    def test_failed_pending_unknown_or_wrong_scope_cannot_export_stale_data(self, sync, rows):
        sync.return_value = None
        with self.assertRaises(LookupError):
            report.export_synced_report("missing")
        for kind, sectors, status in [
            ("ots", [report.NF_SECTOR], "failed"), ("ots", [report.NF_SECTOR], "running"),
            ("mao_de_obra", [report.NF_SECTOR], "completed"),
            ("ots", ["RES-OTHER"], "completed"), ("ots", [report.NF_SECTOR, "RES-OTHER"], "completed"),
        ]:
            sync.return_value = {"tipo": kind, "setores": sectors, "status": status}
            with self.subTest(sync=sync.return_value), self.assertRaises(ValueError):
                report.export_synced_report("job")
        rows.assert_not_called()
